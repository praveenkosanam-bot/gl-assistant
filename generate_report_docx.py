import json
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_word_report(json_path, output_path):
    # Read the JSON data
    with open(json_path, 'r', encoding='utf-8') as f:
        results = json.load(f)

    # Create a new Document
    doc = Document()

    # Title
    title = doc.add_heading('GL Assistant Coverage Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Summary Statistics
    total = len(results)
    counts = {}
    for r in results:
        cls = r.get("classification", "UNKNOWN")
        counts[cls] = counts.get(cls, 0) + 1

    doc.add_heading('Summary Statistics', level=1)
    
    summary_table = doc.add_table(rows=1, cols=3)
    summary_table.style = 'Table Grid'
    hdr_cells = summary_table.rows[0].cells
    hdr_cells[0].text = 'Status'
    hdr_cells[1].text = 'Count'
    hdr_cells[2].text = 'Percentage'

    for cls, count in counts.items():
        row_cells = summary_table.add_row().cells
        row_cells[0].text = cls
        row_cells[1].text = str(count)
        row_cells[2].text = f"{count*100/total:.1f}%"

    doc.add_paragraph("\n")

    # Detailed Results Table
    doc.add_heading('Detailed Test Results', level=1)
    
    # Create the results table
    # Columns: Question, Intent, API Path, Classification, Elapsed (ms)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Question'
    hdr_cells[1].text = 'Intent'
    hdr_cells[2].text = 'API Path'
    hdr_cells[3].text = 'Status'
    hdr_cells[4].text = 'Time (ms)'

    # Adjust column widths
    widths = [Inches(2.5), Inches(1.5), Inches(1.5), Inches(1.0), Inches(0.5)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width

    # Add question rows
    for r in results:
        row_cells = table.add_row().cells
        row_cells[0].text = r.get("question", "")
        row_cells[1].text = r.get("intent", "").replace("balance.", "").replace("journals.", "")
        row_cells[2].text = r.get("api_path", "").replace("/api/db/", "")
        row_cells[3].text = r.get("classification", "")
        row_cells[4].text = str(r.get("elapsed_ms", ""))

        # Optional: styling based on status
        # Note: python-docx doesn't easily support cell background colors without XML manipulation.

    # Save the document
    doc.save(output_path)
    print(f"Word report generated: {output_path}")

if __name__ == "__main__":
    generate_word_report('test_report_full.json', 'test_report.docx')
